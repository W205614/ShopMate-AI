"""Regenerate synthetic AI acceptance files with Pillow, python-docx and ReportLab."""
from pathlib import Path
from io import BytesIO
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.lib.utils import ImageReader

folder = Path(__file__).resolve().parents[1] / "test" / "fixtures" / "ai"
folder.mkdir(parents=True, exist_ok=True)
(folder / "spruce.txt").write_text("合成验收资料，不适用于真实交易。云杉测试保养布建议水温低于30度，不可使用漂白剂。", encoding="utf-8")
(folder / "lamp.md").write_text("# 合成验收资料\n\n晴岚测试台灯的合成保修期限为17天。禁止浸水。不是实际商品承诺。", encoding="utf-8")
doc = Document()
doc.add_heading("碧泉合成资料", 0)
doc.add_paragraph("碧泉测试瓶的容量为420毫升，禁止微波加热。这是合成测试资料，不适用于真实商品。")
doc.save(folder / "bottle.docx")
pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
pdf = canvas.Canvas(str(folder / "text.pdf"), pagesize=(595, 842))
pdf.setFont("STSong-Light", 16)
for n, line in enumerate(["远山合成验收资料", "远山测试杯建议使用35度温水清洁，禁止机洗。", "仅用于隔离验收，不是实际商品承诺。"]):
    pdf.drawString(40, 780 - n * 35, line)
pdf.save()
font_path = Path("C:/Windows/Fonts/msyh.ttc")
if not font_path.exists():
    raise SystemExit("Set a Chinese font path before regenerating OCR fixtures")
image = Image.new("RGB", (1400, 650), "white")
draw = ImageDraw.Draw(image)
font = ImageFont.truetype(str(font_path), 40)
for n, line in enumerate(["晨星合成验收资料", "晨星测试杯：容量420毫升。", "合成质保期限17天。禁止机洗。", "包装破损时不要丢弃包装，先联系人工。", "仅用于隔离验收，非真实商品承诺。"]):
    draw.text((45, 40 + n * 105), line, font=font, fill="black")
for ext in ["png", "jpg", "webp"]:
    image.save(folder / f"image.{ext}")
scan = canvas.Canvas(str(folder / "scan.pdf"), pagesize=(700, 400))
scan.drawImage(ImageReader(image), 0, 30, width=700, height=325)
scan.save()
print("Generated 8 synthetic document and image fixtures")
